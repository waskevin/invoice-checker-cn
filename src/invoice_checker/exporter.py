from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Invoice, Summary


def export_invoices(path: Path, invoices: list[Invoice], summary: Summary) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "核验结果"
    sheet.append(["发票批量核验结果"])
    sheet.append(["文件名", "发票号码", "开票日期", "类型", "销售方", "购买方", "价税合计（小写）", "状态", "说明"])
    for invoice in invoices:
        sheet.append([invoice.file_path.name, invoice.number or "", invoice.issue_date.isoformat() if invoice.issue_date else "",
                      invoice.invoice_kind.value, invoice.seller_name or "", invoice.buyer_name or "",
                      float(invoice.total_with_tax) if invoice.total_with_tax is not None else None,
                      invoice.status.value, invoice.status_message])
    sheet.append([])
    sheet.append(["正常有效张数", summary.valid_count, "有效金额合计", float(summary.valid_total)])
    sheet["A1"].font = Font(bold=True, size=14)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for index, width in enumerate((24, 22, 14, 10, 26, 26, 18, 14, 42), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A3"
    workbook.save(path)

